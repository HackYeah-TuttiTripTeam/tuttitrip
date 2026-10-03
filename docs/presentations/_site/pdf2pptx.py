# /// script
# requires-python = ">=3.11"
# dependencies = ["pymupdf>=1.24", "python-pptx>=1.0"]
# ///
"""PDF deck -> PPTX with real, selectable text (build.mjs runs it after printing the PDF).

Every PDF page becomes a slide of the same size. The page is rendered without its text
as the slide background (photos, cards and icons stay pixel-exact), and every line of
text goes on top as its own text box with the original font, size, colour and links.

Chrome embeds variable fonts as anonymous Type 3 fonts, so the font names come from the
page instead: fonts.json (written by print.mjs) lists the page's text runs with their CSS
font, and each PDF font gets the family and weight most of its text has on the page.

    uv run docs/presentations/_site/pdf2pptx.py deck.pdf deck.pptx [fonts.json]
"""
import io
import json
import sys
from collections import Counter, defaultdict

import pymupdf as fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Pt

DPI_SCALE = 2  # background resolution: 2 px per PDF point


def font_map(pdf, runs):
    """PDF font name -> (family, bold), voted by the page's text runs of the same size."""
    by_size = defaultdict(list)
    for run in runs:
        by_size[round(run['size'] * 2) / 2].append(run)
    votes = defaultdict(Counter)
    for page in pdf:
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                for span in line['spans']:
                    text = ' '.join(span['text'].split())
                    if len(text) < 2:
                        continue
                    size = round(span['size'] * 2) / 2
                    for near in (size, size - 0.5, size + 0.5):
                        for run in by_size.get(near, []):
                            if text in run['text']:
                                votes[span['font']][(run['family'], run['weight'] >= 600)] += len(text)
    return {font: counter.most_common(1)[0][0] for font, counter in votes.items()}


def font_of(span, fonts):
    if span['font'] in fonts:
        return fonts[span['font']]
    if span['font'].startswith('Type3'):
        return None  # unknown: the caller takes a neighbour's font
    name = span['font'].split('+')[-1]
    return name.split('-')[0], bool(span['flags'] & 16)


def line_fonts(spans, fonts):
    """Fonts for a line's spans; spaces and lone characters borrow from a neighbour."""
    found = [font_of(s, fonts) for s in spans]
    known = [f for f in found if f]
    fallback = known[0] if known else ('Atkinson Hyperlegible Next', False)
    out, last = [], None
    for f in found:
        last = f or last
        out.append(f or last or fallback)
    return out


def background(page):
    """The page as a JPEG with all text removed (graphics and images kept)."""
    doc = fitz.open()
    doc.insert_pdf(page.parent, from_page=page.number, to_page=page.number)
    clean = doc[0]
    for block in clean.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            for span in line['spans']:
                clean.add_redact_annot(fitz.Rect(span['bbox']))
    clean.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE)
    pix = clean.get_pixmap(matrix=fitz.Matrix(DPI_SCALE, DPI_SCALE))
    return io.BytesIO(pix.tobytes('jpeg', jpg_quality=88))


def main(src, out, fonts_json=None):
    pdf = fitz.open(src)
    fonts = font_map(pdf, json.load(open(fonts_json, encoding='utf-8'))) if fonts_json else {}
    first = pdf[0].rect
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(first.width), Pt(first.height)
    blank = prs.slide_layouts[6]
    for page in pdf:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(background(page), 0, 0, prs.slide_width, prs.slide_height)
        links = [(fitz.Rect(l['from']), l['uri']) for l in page.get_links() if l.get('uri')]
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                spans = [s for s in line['spans'] if s['text'].strip()]
                if not spans:
                    continue
                x0, y0, x1, y1 = line['bbox']
                # A little slack, no wrapping: a substitute font must not break the line.
                box = slide.shapes.add_textbox(Pt(x0), Pt(y0), Pt((x1 - x0) * 1.08 + 2), Pt(y1 - y0))
                frame = box.text_frame
                frame.word_wrap = False
                frame.auto_size = None
                frame.vertical_anchor = MSO_ANCHOR.TOP
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Emu(0)
                para = frame.paragraphs[0]
                for span, (family, bold) in zip(line['spans'], line_fonts(line['spans'], fonts)):
                    run = para.add_run()
                    run.text = span['text']
                    run.font.name, run.font.bold = family, bold
                    run.font.size = Pt(round(span['size'] * 2) / 2)
                    run.font.color.rgb = RGBColor.from_string(f"{span['color']:06X}")
                    # The link goes on the text box, so the text keeps its colour (PowerPoint
                    # would paint linked text in the theme's blue and underline it).
                    box_rect = fitz.Rect(span['bbox'])
                    for rect, uri in links:
                        if box_rect.get_area() and (box_rect & rect).get_area() > box_rect.get_area() / 2:
                            box.click_action.hyperlink.address = uri
    prs.save(out)
    print(f'pdf2pptx: {len(pdf)} slide(s) -> {out}')


if __name__ == '__main__':
    main(*sys.argv[1:4])
